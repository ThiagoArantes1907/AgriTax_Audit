#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
================================================================================
AgriTax Audit — MÓDULO FRIGORÍFICO (Fase 1)
Revisão tributária automatizada de frigorífico de abate de bovinos (Lucro Real)
por cruzamento EFD-Contribuições × EFD-ICMS/IPI.
--------------------------------------------------------------------------------
Uso:
    python modulo_frigorifico.py <pasta_com_arquivos_txt> [--saida decisoria.xlsx]
                                 [--per situacao_per.csv]

  A pasta deve conter os .txt das EFDs (Contribuições e ICMS/IPI) do período.
  O módulo detecta o tipo e a competência de cada arquivo automaticamente,
  pareia por mês e roda os testes A1-A7 e B1-B4 (matriz da seção 9.2 do
  material de revisão). A saída é a PLANILHA DECISÓRIA (padrão AgriTax).

REGRA DE PERÍODOS FECHADOS (PER/DCOMP de ressarcimento):
  O trabalho recebe todos os pedidos de ressarcimento já efetuados + consulta
  de situação. Arquivo --per (CSV "competencia;situacao", ex.: "05/2026;PER
  DEFERIDA"). Período cujo ressarcimento esteja DEFERIDA / ANÁLISE CONCLUÍDA /
  HOMOLOGADO é tratado como FECHADO:
    * OPORTUNIDADES do período fechado são todas mapeadas e trazidas; a ação
      de APROPRIAÇÃO EXTEMPORÂNEA no PRIMEIRO período ainda aberto (coluna
      "Competência de destino") aplica-se SOMENTE aos benefícios inseríveis
      mediante registros dos BLOCOS A, C, D e F da EFD-Contribuições
      (documentos: serviços/A, NF-e/C, transportes/D, demais operações e
      imobilizado/F). Benefício que só se materializa como ajuste de
      apuração (bloco M) não é inserível em período de destino — é trazido
      com a ação "SEM INSERÇÃO EM PERÍODO ABERTO", para avaliação de via
      alternativa.
    * RISCOS identificados em período fechado NÃO viram ação de correção:
      compõem a aba MATRIZ DE RISCO (exposição, status do PER, ação de
      gestão). Riscos de períodos abertos seguem em Expurgos-Riscos.
  Sem o arquivo --per, todos os períodos são tratados como abertos.

Testes implementados (Fase 1):
  A1  Índice de exportação (§7º art. 33 Lei 12.058/2009), mês a mês
  A2  Crédito presumido art. 33 não apropriado / insuficiente (quantificado)
  A3  Presumido com alíquota fora da vigência (50% / 45% pós-LC 224)
  A4  Crédito básico indevido em gado de PJ (suspensão art. 32) — RISCO
  A5  Rateio contaminado por NCM fora da lista (couro cap.41, tripas 05.04)
  A6  Venda de carne p/ CNPJ sem suspensão (art. 32, II) — ação condicional
  A7  Competência >= abr/2026 com carnes internas sem débito de 10% (LC 224)
  B1  NF-e de entrada na EFD-ICMS ausente das Contribuições (por classe)
  B2  CT-e tomado (D100) sem crédito nas Contribuições
  B3  Classe de insumo inteira sem crédito no período
  B4  Base de crédito sem expurgo do ICMS (competências >= mai/2023) — RISCO
  B5  Bloco M vazio com documentos / Contribuições sem movimento com ICMS ativo

AJUSTES APRENDIDOS EM ARQUIVO REAL (caso Boi Bravo, 09/2026):
  * Regime lido do 0110 (COD_INC_TRIB): em competência CUMULATIVA (2) não há
    direito a crédito básico — os achados B1/B2/B3 são trazidos como SIMULAÇÃO
    (quanto valeria a não cumulatividade), nunca como ação CREDITAR.
  * Receita consolidada em F550 (regime cumulativo) conta como escrituração:
    não dispara B5 "sem movimento".
  * Base do presumido (A2) só com CFOP de AQUISIÇÃO do gado (1101/1102/2101/
    2102...); entradas por encomenda (1901/1924), remessa e imobilizado ficam fora.
  * Classe de insumo por PALAVRA-CHAVE só quando o NCM é desconhecido ou está
    fora dos capítulos 01-24 (alimentos/animais) e 41 (couro) — "(430 CAIXA)" na
    descrição de miúdos bovinos NÃO é embalagem.
  * --desde MM/AAAA limita o escopo temporal (prescrição corre mês a mês).
  * MATRIZ + FILIAIS (caso Frigorífico Modelo): a EFD-ICMS é por estabelecimento
    e a Contribuições é centralizada — EfdIcmsConsolidada soma os arquivos
    vigentes de cada CNPJ da competência, prefixando cod_item/cod_part pelo
    CNPJ (as numerações são independentes por estabelecimento).
  * SEM EFD-CONTRIBUIÇÕES no acervo: o motor roda em modo "universo ICMS"
    (--sem-contrib ou automático quando a competência só tem ICMS) — B1/B2
    listam TODO o universo creditável, A1/A2 saem normalmente e a coluna
    Apuração marca "SEM EFD-CONTRIBUIÇÕES" (pedir ao titular).
  * Energia elétrica também entra por C100/C170 (NCM 2716, CFOP 1252/1253) —
    classe "Energia elétrica" (art. 3º, III/IX). Combustíveis (2710) entram
    como classe condicionada a prova (uso no processo × frota de entrega).
  * NCM "dummy" (05119999, 00000000, vazio) é tratado como desconhecido —
    ERPs usam 0511.99.99 como curinga (Modelo: "IMOBILIZADO-MAQUINAS").

IMPORTANTE (prática da casa): validar as posições de campo dos leiautes contra
arquivos reais do cliente antes do uso em produção — os índices abaixo seguem
o Guia Prático (EFD-Contribuições e EFD-ICMS/IPI) e têm leitura defensiva.
================================================================================
"""

import os
import re
import sys
import argparse
from collections import defaultdict
from datetime import date

# ============================================================================
# 1. DICIONÁRIO SETORIAL DO FRIGORÍFICO (bovinos — Lei 12.058/2009)
# ============================================================================

NCM_GADO = ("0102",)                      # boi vivo
NCM_CARNES_IN_NATURA = ("0201", "0202")   # carnes bovinas frescas/congeladas

# Lista EXAUSTIVA do caput do art. 33 da Lei 12.058/2009 (bovinos) — prefixos
NCM_LISTA_ART33 = (
    "0201", "0202", "02061000", "020620", "020621", "020629",
    "02102000", "05069000", "05100010", "150210",
)
# Fora da lista: contaminam o numerador do rateio (§7º) se exportados
NCM_FORA_RATEIO = ("41", "0504")          # couro (cap. 41), tripas (05.04)

# Classes de insumo típicas do frigorífico (teste B1/B3)
CLASSES_INSUMO = {
    "Embalagens":          {"ncm": ("3920", "3921", "3923", "4819", "4821", "7310"),
                            "kw": ("EMBALAG", "FILME", "BANDEJA", "CAIXA DE PAP", "CAIXA PAPEL", "CX PAPEL",
                                   "CAIXA PLAST", "ROTUL", "ETIQUET", "STRETCH")},
    "Gases de processo":   {"ncm": ("2814", "2811", "2804"),
                            "kw": ("AMONIA", "CO2", "CARBONICO", "NITROGENIO")},
    "Sanitizantes/limpeza": {"ncm": ("3402", "2828", "380894", "2915"),
                            "kw": ("DETERGENTE", "SANITIZ", "HIPOCLORITO", "PERACETICO", "SODA CAUSTICA", "DESINFET")},
    "EPI/uniformes":       {"ncm": ("6116", "6210", "6401", "6403", "392620", "6505"),
                            "kw": ("LUVA", "BOTA", "AVENTAL", "TOUCA", "UNIFORME", "EPI", "PROTETOR")},
    "Lenha/combustível caldeira": {"ncm": ("4401", "2701"),
                            "kw": ("LENHA", "CAVACO", "BRIQUETE", "BIOMASSA")},
    "Peças/manutenção":    {"ncm": ("8208", "8438", "7318", "8482", "4010"),
                            "kw": ("ROLAMENTO", "CORREIA", "SERRA", "FACA", "MANUTEN", "PECA ", "PEÇA")},
    "Laboratório/efluentes": {"ncm": (),
                            "kw": ("ANALISE", "ANÁLISE", "LABORATOR", "EFLUENTE", " ETE ", "MICROBIOL")},
    "Energia elétrica":    {"ncm": ("2716",),
                            "kw": ("ENERGIA ELETRICA", "ENERGIA ELÉTRICA", "FORNECIMENTO DE ENERGIA")},
    "Combustíveis (prova de uso)": {"ncm": ("2710", "3403", "2711"),
                            "kw": ("OLEO DIESEL", "ÓLEO DIESEL", "DIESEL S10", "GLP", "LUBRIFICANTE")},
}

# CFOPs — receitas e operações
CFOP_EXPORT_DIRETA_PREF = "7"                     # 7xxx
CFOP_FIM_ESPECIFICO = {"5501", "5502", "6501", "6502"}
CFOP_TRANSFER = {"5151", "5152", "5155", "5156", "6151", "6152", "6155", "6156"}
CFOP_DEVOLUCAO_PREF = ("5201", "5202", "5410", "5411", "6201", "6202", "6410", "6411", "1", "2")
CFOP_VENDA_INTERNA_PREF = ("51", "52", "53", "54", "55", "61", "62", "63", "64", "65")

# CSTs PIS/COFINS
CST_CRED_BASICO = {"50", "51", "52", "53", "54", "55", "56"}
CST_CRED_PRESUMIDO = {"60", "61", "62", "63", "64", "65", "66"}
CST_SAIDA_ZERO = "06"
CST_SAIDA_SUSPENSAO = "09"
COD_CRED_PRESUMIDO = {"106", "206", "306"}        # Tab. 4.3.6 — presumido agroindústria

# Vigências legislativas (motor por competência)
COMP_LC224 = (2026, 4)          # efeitos PIS/COFINS da LC 224/2025
COMP_EXPURGO_ICMS = (2023, 5)   # Lei 14.592/2023 — ICMS fora da base do crédito
ALIQ_BASICA = (0.0165, 0.076)

# Situações de PER/DCOMP que FECHAM o período para inserção/retificação
_STATUS_FECHADO_KW = ("DEFERID", "HOMOLOGAD", "ANALISE CONCLUIDA")


def _norm_txt(s):
    tab = str.maketrans("ÁÀÂÃÉÊÍÓÔÕÚÜÇáàâãéêíóôõúüç", "AAAAEEIOOOUUCaaaaeeiooouuc")
    return (s or "").translate(tab).upper().strip()


def periodo_fechado(status):
    n = _norm_txt(status)
    return any(k in n for k in _STATUS_FECHADO_KW)


def _comps_de_texto(txt):
    """Interpreta competência: 'MM/AAAA', data, ou trimestre ('1T2024',
    '1º TRI 2024', '1º TRIMESTRE/2024') → lista de (ano, mes)."""
    if hasattr(txt, "year"):                      # datetime/date de célula
        return [(txt.year, txt.month)]
    s = _norm_txt(str(txt))
    m = re.match(r"(\d{1,2})\s*[/\-.]\s*(\d{4})", s)
    if m and 1 <= int(m.group(1)) <= 12:
        return [(int(m.group(2)), int(m.group(1)))]
    m = re.match(r"(\d{4})\s*[/\-.]\s*(\d{1,2})$", s)  # AAAA/MM
    if m and 1 <= int(m.group(2)) <= 12:
        return [(int(m.group(1)), int(m.group(2)))]
    m = re.search(r"([1-4])\s*[ºO]?\s*(T|TRI|TRIM|TRIMESTRE)\w*\s*[/\-.]?\s*(\d{4})", s)
    if m:
        tri, ano = int(m.group(1)), int(m.group(3))
        return [(ano, (tri - 1) * 3 + k) for k in (1, 2, 3)]
    return []


def carregar_situacao_per(path):
    """Situação dos PERs de ressarcimento por competência.
    Aceita: CSV 'competencia;situacao' OU planilha .xlsx/.xlsm (detecta as
    colunas de competência/período e de situação/status pelo cabeçalho).
    Competências: MM/AAAA, data ou trimestre ('1º TRI 2024')."""
    mapa = {}
    if not path or not os.path.isfile(path):
        return mapa
    if path.lower().endswith((".xlsx", ".xlsm")):
        try:
            from openpyxl import load_workbook
        except ImportError:
            print("AVISO: openpyxl indisponível para ler %s" % path)
            return mapa
        wb = load_workbook(path, data_only=True)
        for ws in wb.worksheets:
            col_comp = col_sit = None
            cab_row = None
            for row in ws.iter_rows(min_row=1, max_row=10):
                for cel in row:
                    t = _norm_txt(str(cel.value)) if cel.value is not None else ""
                    if col_comp is None and any(k in t for k in ("COMPET", "PERIODO", "PER. APUR", "TRIMESTRE")):
                        col_comp, cab_row = cel.column, cel.row
                    if col_sit is None and any(k in t for k in ("SITUAC", "STATUS")):
                        col_sit = cel.column
                if col_comp and col_sit:
                    break
            if not (col_comp and col_sit):
                continue
            for row in ws.iter_rows(min_row=(cab_row or 1) + 1):
                vc = row[col_comp - 1].value if col_comp - 1 < len(row) else None
                vs = row[col_sit - 1].value if col_sit - 1 < len(row) else None
                if vc is None or vs is None:
                    continue
                for comp in _comps_de_texto(vc):
                    atual = mapa.get(comp)
                    # se já houver situação p/ a competência, a ABERTA prevalece
                    # (basta um PER não decidido para o período não estar fechado)
                    if atual is None or (periodo_fechado(atual) and not periodo_fechado(str(vs))):
                        mapa[comp] = str(vs).strip()
        return mapa
    with open(path, "r", encoding="utf-8", errors="replace") as fh:
        for linha in fh:
            linha = linha.strip()
            if not linha or _norm_txt(linha).startswith("COMPET"):
                continue
            partes = re.split(r"[;,\t]", linha)
            if len(partes) < 2:
                continue
            for comp in _comps_de_texto(partes[0].strip()):
                mapa[comp] = partes[1].strip()
    return mapa


def aliq_presumido(comp):
    """Alíquotas do crédito presumido art. 33 (PIS, COFINS) por competência."""
    if comp >= COMP_LC224:                       # redução de 10% (LC 224/2025)
        return (0.0165 * 0.45, 0.076 * 0.45)     # 0,7425% / 3,42%
    return (0.0165 * 0.50, 0.076 * 0.50)         # 0,825% / 3,80%


def aliq_carne_interna(comp):
    """Débito sobre venda interna de carnes: zero até mar/2026; 10% da alíquota depois."""
    if comp >= COMP_LC224:
        return (0.0165 * 0.10, 0.076 * 0.10)     # 0,165% / 0,76%
    return (0.0, 0.0)


# ============================================================================
# 2. PARSERS EFD (leitura defensiva, latin-1, pipe-delimited)
# ============================================================================

def _f(campos, i, default=""):
    return campos[i].strip() if i < len(campos) else default


def _num(s):
    s = (s or "").strip()
    if not s:
        return 0.0
    try:
        return float(s.replace(".", "").replace(",", "."))
    except ValueError:
        return 0.0


def _data(s):
    s = (s or "").strip()
    if len(s) == 8 and s.isdigit():
        try:
            return date(int(s[4:8]), int(s[2:4]), int(s[0:2]))
        except ValueError:
            return None
    return None


def _split(linha):
    linha = linha.rstrip("\r\n")
    if not linha.startswith("|"):
        return None
    return [c for c in linha[1:].rstrip("|").split("|")]  # campos[0] = REG


def detectar_tipo(path):
    """Retorna (tipo, competencia, cnpj, prioridade). Prioridade ordena versões da
    mesma competência: (retificadora/substituto=1, timestamp do nome ou mtime)."""
    nome = os.path.basename(path).upper()
    with open(path, "r", encoding="latin-1", errors="replace") as fh:
        for linha in fh:
            c = _split(linha)
            if not c or c[0] != "0000":
                continue
            if _data(_f(c, 3)):                       # EFD-ICMS
                d = _data(_f(c, 3))
                cod_fin = _f(c, 2)                    # 0=original, 1=substituto
                m = re.search(r"(\d{8})-(\d{8})-(\d)-", nome)
                prio = (1 if cod_fin == "1" else 0, os.path.getmtime(path))
                return "ICMS", (d.year, d.month), re.sub(r"\D", "", _f(c, 6)), prio
            if _data(_f(c, 5)):                       # EFD-Contribuições
                d = _data(_f(c, 5))
                retif = 1 if "RETIFICADORA" in nome else 0
                m = re.search(r"_(\d{14})_", nome)    # timestamp de transmissão
                ts = float(m.group(1)) if m else os.path.getmtime(path)
                return "CONTRIB", (d.year, d.month), re.sub(r"\D", "", _f(c, 8)), (retif, ts)
            break
    return None, None, None, None


class EfdIcms:
    """Universo real de operações: participantes, itens/NCM, docs, itens, CT-e."""

    def __init__(self, path):
        self.path = path
        self.participantes = {}     # cod -> {"cnpj","cpf","nome"}
        self.itens = {}             # cod_item -> {"ncm","descr"}
        self.docs = []              # C100 + filhos C170
        self.ctes = []              # D100
        self.c190 = []              # (ind_oper_doc?, cst, cfop, vl_opr) — via docs
        self._parse()

    def _parse(self):
        doc = None
        with open(self.path, "r", encoding="latin-1", errors="replace") as fh:
            for linha in fh:
                c = _split(linha)
                if not c:
                    continue
                reg = c[0]
                if reg == "0150":
                    self.participantes[_f(c, 1)] = {
                        "nome": _f(c, 2),
                        "cnpj": re.sub(r"\D", "", _f(c, 4)),
                        "cpf": re.sub(r"\D", "", _f(c, 5)),
                    }
                elif reg == "0200":
                    self.itens[_f(c, 1)] = {"descr": _f(c, 2).upper(),
                                            "ncm": re.sub(r"\D", "", _f(c, 7))}
                elif reg == "C100":
                    doc = {"ind_oper": _f(c, 1), "cod_part": _f(c, 3),
                           "cod_sit": _f(c, 5), "num": _f(c, 7),
                           "chave": re.sub(r"\D", "", _f(c, 8)),
                           "vl_doc": _num(_f(c, 11)), "vl_merc": _num(_f(c, 15)),
                           "vl_icms": _num(_f(c, 21)), "itens": []}
                    if doc["cod_sit"] in ("00", "01", "06", "07", "08"):
                        self.docs.append(doc)
                    else:
                        doc = None  # cancelados/denegados fora
                elif reg == "C170" and doc is not None:
                    doc["itens"].append({
                        "cod_item": _f(c, 2), "vl_item": _num(_f(c, 6)),
                        "cfop": _f(c, 10), "vl_icms": _num(_f(c, 14)),
                    })
                elif reg == "C190" and doc is not None:
                    self.c190.append({"ind_oper": doc["ind_oper"],
                                      "cfop": _f(c, 2), "vl_opr": _num(_f(c, 4))})
                elif reg == "D100":
                    self.ctes.append({"ind_oper": _f(c, 1), "cod_part": _f(c, 3),
                                      "chave": re.sub(r"\D", "", _f(c, 9)),
                                      "vl_serv": _num(_f(c, 17))})

    # helpers ----------------------------------------------------------------
    def ncm_do_item(self, cod_item):
        return self.itens.get(cod_item, {}).get("ncm", "")

    def descr_do_item(self, cod_item):
        return self.itens.get(cod_item, {}).get("descr", "")

    def fornecedor(self, cod_part):
        return self.participantes.get(cod_part, {"cnpj": "", "cpf": "", "nome": ""})


class EfdIcmsConsolidada:
    """Matriz + filiais numa competência: soma os EfdIcms de cada CNPJ, prefixando
    cod_item/cod_part pelo CNPJ (numerações independentes por estabelecimento)."""

    def __init__(self, paths):
        self.paths = list(paths)
        self.participantes, self.itens = {}, {}
        self.docs, self.ctes, self.c190 = [], [], []
        self.cnpjs = []
        for p in self.paths:
            e = EfdIcms(p)
            cnpj = ""
            with open(p, "r", encoding="latin-1", errors="replace") as fh:
                for linha in fh:
                    c = _split(linha)
                    if c and c[0] == "0000":
                        cnpj = re.sub(r"\D", "", _f(c, 6))
                        break
            pref = (cnpj or str(len(self.cnpjs))) + ":"
            self.cnpjs.append(cnpj)
            for k, v in e.participantes.items():
                self.participantes[pref + k] = v
            for k, v in e.itens.items():
                self.itens[pref + k] = v
            for d in e.docs:
                d = dict(d)
                d["cod_part"] = pref + d["cod_part"]
                d["estab"] = cnpj
                d["itens"] = [dict(it, cod_item=pref + it["cod_item"]) for it in d["itens"]]
                self.docs.append(d)
            for ct in e.ctes:
                ct = dict(ct)
                ct["cod_part"] = pref + ct["cod_part"]
                ct["estab"] = cnpj
                self.ctes.append(ct)
            self.c190.extend(e.c190)

    def ncm_do_item(self, cod_item):
        return self.itens.get(cod_item, {}).get("ncm", "")

    def descr_do_item(self, cod_item):
        return self.itens.get(cod_item, {}).get("descr", "")

    def fornecedor(self, cod_part):
        return self.participantes.get(cod_part, {"cnpj": "", "cpf": "", "nome": ""})


class EfdContribVazia:
    """Stub para competência SEM EFD-Contribuições no acervo (modo universo ICMS)."""
    path = None
    chaves = set()
    itens = {}
    participantes = {}
    creditos = []
    det_cred = []
    linhas_item = []
    regime = ""
    receita_f550 = 0.0
    ausente = True

    def ncm_do_item(self, cod_item):
        return ""


class EfdContrib:
    """Apuração declarada: chaves, créditos M100/M500, M105/M505, itens C170."""
    ausente = False

    def __init__(self, path):
        self.path = path
        self.chaves = set()          # C100 + D100 escriturados
        self.itens = {}              # 0200
        self.participantes = {}      # 0150
        self.creditos = []           # M100/M500: contribuição, cod_cred, bc, vl_cred
        self.det_cred = []           # M105/M505: contribuição, cst, bc
        self.linhas_item = []        # C170 com CST/valores PIS-COFINS
        self.regime = ""             # 0110 COD_INC_TRIB: 1 não cumulativo, 2 cumulativo, 3 ambos
        self.receita_f550 = 0.0      # receitas consolidadas (regime cumulativo)
        self._parse()

    def _parse(self):
        doc = None
        with open(self.path, "r", encoding="latin-1", errors="replace") as fh:
            for linha in fh:
                c = _split(linha)
                if not c:
                    continue
                reg = c[0]
                if reg == "0110":
                    self.regime = _f(c, 1)
                elif reg == "F550":
                    self.receita_f550 += _num(_f(c, 1))
                elif reg == "0150":
                    self.participantes[_f(c, 1)] = {
                        "nome": _f(c, 2),
                        "cnpj": re.sub(r"\D", "", _f(c, 4)),
                        "cpf": re.sub(r"\D", "", _f(c, 5))}
                elif reg == "0200":
                    self.itens[_f(c, 1)] = {"descr": _f(c, 2).upper(),
                                            "ncm": re.sub(r"\D", "", _f(c, 7))}
                elif reg == "C100":
                    doc = {"ind_oper": _f(c, 1), "cod_part": _f(c, 3),
                           "chave": re.sub(r"\D", "", _f(c, 8)),
                           "vl_doc": _num(_f(c, 11)), "vl_icms": _num(_f(c, 21))}
                    if doc["chave"]:
                        self.chaves.add(doc["chave"])
                elif reg == "C170" and doc is not None:
                    self.linhas_item.append({
                        "ind_oper": doc["ind_oper"], "cod_part": doc["cod_part"],
                        "chave": doc["chave"], "vl_icms_doc": doc["vl_icms"],
                        "cod_item": _f(c, 2), "vl_item": _num(_f(c, 6)),
                        "cfop": _f(c, 10),
                        "cst_pis": _f(c, 15).zfill(2)[:2] if _f(c, 15) else "",
                        "bc_pis": _num(_f(c, 16)), "vl_pis": _num(_f(c, 20)),
                        "cst_cof": _f(c, 21).zfill(2)[:2] if _f(c, 21) else "",
                        "bc_cof": _num(_f(c, 22)), "vl_cof": _num(_f(c, 26)),
                    })
                elif reg == "D100":
                    ch = re.sub(r"\D", "", _f(c, 9))
                    if ch:
                        self.chaves.add(ch)
                elif reg in ("M100", "M500"):
                    self.creditos.append({
                        "trib": "PIS" if reg == "M100" else "COFINS",
                        "cod_cred": _f(c, 1), "bc": _num(_f(c, 3)),
                        "aliq": _num(_f(c, 4)), "vl_cred": _num(_f(c, 7)),
                    })
                elif reg in ("M105", "M505"):
                    self.det_cred.append({
                        "trib": "PIS" if reg == "M105" else "COFINS",
                        "cst": _f(c, 2).zfill(2)[:2] if _f(c, 2) else "",
                        "bc": _num(_f(c, 6)) or _num(_f(c, 3)),
                    })

    def ncm_do_item(self, cod_item):
        return self.itens.get(cod_item, {}).get("ncm", "")


# ============================================================================
# 3. MOTOR DE TESTES
# ============================================================================

NCM_CAP_SEM_KW = tuple("%02d" % i for i in range(1, 25) if i != 5) + ("41",)   # alimentos/animais/couro
NCM_DUMMY = {"", "05119999", "00000000", "00"}    # curingas de ERP — tratar como desconhecido


def _classe_insumo(ncm, descr):
    if ncm in NCM_DUMMY:
        ncm = ""
    for nome, cfg in CLASSES_INSUMO.items():
        if ncm and any(ncm.startswith(p) for p in cfg["ncm"]):
            return nome
    # palavra-chave só quando o NCM não resolve E não é produto alimentar/animal
    if ncm and ncm[:2] in NCM_CAP_SEM_KW:
        return None
    for nome, cfg in CLASSES_INSUMO.items():
        if descr and any(k in descr for k in cfg["kw"]):
            return nome
    return None


# CFOPs de AQUISIÇÃO (compra p/ industrialização/comercialização) — base do presumido
CFOP_AQUISICAO = {"1101", "1102", "1111", "1113", "1116", "1117", "1118", "1120",
                  "1121", "1122", "2101", "2102", "2111", "2113", "2116", "2117",
                  "2118", "2120", "2121", "2122", "3101", "3102"}


def _eh_aquisicao(cfop):
    return cfop in CFOP_AQUISICAO


def _eh_receita(cfop):
    if not cfop:
        return False
    if cfop in CFOP_TRANSFER:
        return False
    if any(cfop.startswith(p) for p in CFOP_DEVOLUCAO_PREF):
        return False
    return cfop.startswith(CFOP_EXPORT_DIRETA_PREF) or \
        any(cfop.startswith(p) for p in CFOP_VENDA_INTERNA_PREF)


def _eh_export(cfop):
    return cfop.startswith(CFOP_EXPORT_DIRETA_PREF) or cfop in CFOP_FIM_ESPECIFICO


class ResultadoMensal:
    def __init__(self, comp):
        self.comp = comp
        self.receita_total = 0.0
        self.receita_export = 0.0
        self.receita_export_lista = 0.0
        self.receita_export_fora = 0.0        # A5
        self.base_gado_pf = 0.0
        self.base_gado_pj = 0.0
        self.presumido_escriturado = 0.0
        self.presumido_devido = 0.0
        self.risco_a4 = []                    # gado PJ com crédito básico
        self.risco_a7 = 0.0                   # base carnes internas CST 06 pós-LC224
        self.a6_por_adquirente = defaultdict(float)
        self.b1 = []                          # notas descobertas no cruzamento
        self.b2 = []                          # CT-e sem crédito
        self.b3_classes = {}                  # classe -> (existe_icms, tem_cred_contrib)
        self.risco_b4 = []                    # base sem expurgo de ICMS
        self.a3_alertas = []
        # Situação PER/DCOMP do período (regra de períodos fechados)
        self.fonte_receita = ""
        self.sem_apuracao = False             # B5: docs escriturados, bloco M vazio
        self.contrib_vazia = False            # B5: contrib sem movimento c/ ICMS ativo
        self.per_status = "SEM PER / ABERTO"
        self.regime = ""                      # 0110 da Contribuições
        self.cumulativo = False               # sem direito a crédito básico
        self.contrib_ausente = False          # competência sem EFD-Contribuições no acervo
        self.estabs = []                      # CNPJs consolidados (matriz+filiais)
        self.fechado = False
        self.destino = None                   # (ano,mes) p/ apropriação extemporânea


def revisar_competencia(comp, icms, contrib):
    r = ResultadoMensal(comp)

    # ---------------- A1: índice de exportação ------------------------------
    # Fonte preferencial: C170 de SAÍDAS da EFD-Contribuições (tem NCM+CST).
    # Fallback 1: C170 de saídas da EFD-ICMS. Fallback 2: C190 (sem NCM).
    rec_contrib = [li for li in contrib.linhas_item
                   if li["ind_oper"] == "1" and _eh_receita(li["cfop"])]
    if rec_contrib:
        r.fonte_receita = "EFD-Contribuicoes (C170 saidas)"
        for li in rec_contrib:
            r.receita_total += li["vl_item"]
            if _eh_export(li["cfop"]):
                r.receita_export += li["vl_item"]
                ncm = contrib.ncm_do_item(li["cod_item"])
                if any(ncm.startswith(p) for p in NCM_LISTA_ART33):
                    r.receita_export_lista += li["vl_item"]
                elif any(ncm.startswith(p) for p in NCM_FORA_RATEIO):
                    r.receita_export_fora += li["vl_item"]      # A5
    else:
        r.fonte_receita = "EFD-ICMS (C170 saidas)"
        for doc in icms.docs:
            if doc["ind_oper"] != "1":
                continue
            for it in doc["itens"]:
                cfop = it["cfop"]
                if not _eh_receita(cfop):
                    continue
                r.receita_total += it["vl_item"]
                if _eh_export(cfop):
                    r.receita_export += it["vl_item"]
                    ncm = icms.ncm_do_item(it["cod_item"])
                    if any(ncm.startswith(p) for p in NCM_LISTA_ART33):
                        r.receita_export_lista += it["vl_item"]
                    elif any(ncm.startswith(p) for p in NCM_FORA_RATEIO):
                        r.receita_export_fora += it["vl_item"]  # A5
    if r.receita_total == 0:
        # saídas de emissão própria sem C170: usar C190 (agregado por CFOP)
        for x in icms.c190:
            if x["ind_oper"] == "1" and _eh_receita(x["cfop"]):
                r.receita_total += x["vl_opr"]
                if _eh_export(x["cfop"]):
                    r.receita_export += x["vl_opr"]
                    r.receita_export_lista += x["vl_opr"]   # sem NCM: aproximação
        r.fonte_receita = "EFD-ICMS C190 (agregado, sem NCM)"
    indice = (r.receita_export_lista / r.receita_total) if r.receita_total else 0.0

    # ---------------- A2/A4/A6: entradas de gado e saídas de carne ----------
    for doc in icms.docs:
        forn = icms.fornecedor(doc["cod_part"])
        for it in doc["itens"]:
            ncm = icms.ncm_do_item(it["cod_item"])
            if doc["ind_oper"] == "0" and any(ncm.startswith(p) for p in NCM_GADO) \
                    and _eh_aquisicao(it["cfop"]):
                if forn.get("cpf"):
                    r.base_gado_pf += it["vl_item"]
                else:
                    r.base_gado_pj += it["vl_item"]
            if doc["ind_oper"] == "1" and any(ncm.startswith(p) for p in NCM_CARNES_IN_NATURA):
                if forn is not None and icms.fornecedor(doc["cod_part"]).get("cnpj"):
                    if not _eh_export(it["cfop"]):
                        r.a6_por_adquirente[icms.fornecedor(doc["cod_part"])["nome"] or
                                            icms.fornecedor(doc["cod_part"])["cnpj"]] += it["vl_item"]

    ap, ac = aliq_presumido(comp)
    base_gado = r.base_gado_pf + r.base_gado_pj      # PJ agropec. c/ suspensão também gera (SC 80/2021)
    r.presumido_devido = base_gado * (ap + ac) * indice
    r.presumido_escriturado = sum(cr["vl_cred"] for cr in contrib.creditos
                                  if cr["cod_cred"] in COD_CRED_PRESUMIDO)
    if r.presumido_escriturado == 0:
        r.presumido_escriturado = 0.0
        # fallback: bases M105/M505 com CST presumido × alíquota da vigência
        bc61 = sum(d["bc"] for d in contrib.det_cred if d["cst"] in CST_CRED_PRESUMIDO)
        if bc61:
            r.presumido_escriturado = bc61 * (ap + ac) / 2.0  # aproximação conservadora

    # ---------------- A3: alíquota efetiva do presumido ---------------------
    for cr in contrib.creditos:
        if cr["cod_cred"] in COD_CRED_PRESUMIDO and cr["bc"] > 0:
            esperada = ap if cr["trib"] == "PIS" else ac
            efetiva = cr["vl_cred"] / cr["bc"]
            if esperada and abs(efetiva - esperada) / esperada > 0.05:
                r.a3_alertas.append((cr["trib"], efetiva, esperada))

    # ---------------- A4 / A7 / B4 (sobre a EFD-Contribuições) --------------
    for li in contrib.linhas_item:
        ncm = contrib.ncm_do_item(li["cod_item"])
        if li["ind_oper"] == "0":
            part = contrib.participantes.get(li["cod_part"], {})
            if any(ncm.startswith(p) for p in NCM_GADO) and part.get("cnpj") \
                    and li["cst_pis"] in CST_CRED_BASICO and (li["vl_pis"] + li["vl_cof"]) > 0:
                r.risco_a4.append({"chave": li["chave"], "valor": li["vl_item"],
                                   "credito": li["vl_pis"] + li["vl_cof"]})
            if comp >= COMP_EXPURGO_ICMS and li["cst_pis"] in CST_CRED_BASICO \
                    and li["vl_icms_doc"] > 0 and li["bc_pis"] > 0:
                # base declarada maior que (item - icms proporcional)?
                if li["bc_pis"] > li["vl_item"] - (li["vl_icms_doc"] * 0.0):
                    # teste fino exige ICMS por item; nível doc: sinaliza quando bc==vl_item
                    if abs(li["bc_pis"] - li["vl_item"]) < 0.01 and li["vl_icms_doc"] > 0:
                        r.risco_b4.append({"chave": li["chave"], "bc": li["bc_pis"],
                                           "vl_item": li["vl_item"]})
        else:
            if comp >= COMP_LC224 and any(ncm.startswith(p) for p in NCM_CARNES_IN_NATURA) \
                    and li["cst_pis"] == CST_SAIDA_ZERO and not _eh_export(li["cfop"]):
                r.risco_a7 += li["vl_item"]

    # ---------------- B1/B2/B3: cruzamento ICMS × Contribuições -------------
    tem_cred_classe = defaultdict(bool)
    for li in contrib.linhas_item:
        if li["ind_oper"] == "0" and li["cst_pis"] in CST_CRED_BASICO:
            cl = _classe_insumo(contrib.ncm_do_item(li["cod_item"]), "")
            if cl:
                tem_cred_classe[cl] = True

    existe_icms_classe = defaultdict(float)
    for doc in icms.docs:
        if doc["ind_oper"] != "0":
            continue
        ausente = doc["chave"] and doc["chave"] not in contrib.chaves
        for it in doc["itens"]:
            ncm = icms.ncm_do_item(it["cod_item"])
            descr = icms.descr_do_item(it["cod_item"])
            cl = _classe_insumo(ncm, descr)
            if cl:
                existe_icms_classe[cl] += it["vl_item"]
                if ausente:
                    base = max(it["vl_item"] - it["vl_icms"], 0.0)
                    r.b1.append({"classe": cl, "chave": doc["chave"],
                                 "fornecedor": icms.fornecedor(doc["cod_part"])["nome"]
                                 + (" [est. %s]" % doc["estab"][-4:] if doc.get("estab") else ""),
                                 "valor": it["vl_item"],
                                 "credito": base * sum(ALIQ_BASICA)})
    for cte in icms.ctes:
        if cte["ind_oper"] == "0" and cte["chave"] and cte["chave"] not in contrib.chaves:
            r.b2.append({"chave": cte["chave"], "vl_serv": cte["vl_serv"],
                         "credito": cte["vl_serv"] * sum(ALIQ_BASICA)})
    classes_b1 = {x["classe"] for x in r.b1}
    for cl, total in existe_icms_classe.items():
        # B3 só quando a classe não tem crédito escriturado NEM notas já
        # descobertas no B1 (evita dupla contagem do mesmo achado)
        r.b3_classes[cl] = (total, tem_cred_classe.get(cl, False) or cl in classes_b1)

    # ---------------- A8/B5: consistência estrutural ------------------------
    r.regime = contrib.regime
    r.cumulativo = contrib.regime == "2"
    r.contrib_ausente = bool(getattr(contrib, "ausente", False))
    r.estabs = list(getattr(icms, "cnpjs", []))
    tem_docs_contrib = bool(contrib.linhas_item) or bool(contrib.chaves)
    tem_apuracao = bool(contrib.creditos) or bool(contrib.det_cred)
    # cumulativo não tem bloco M de créditos: B5 "sem apuração" só faz sentido no não cumulativo
    r.sem_apuracao = tem_docs_contrib and not tem_apuracao and not r.cumulativo
    r.contrib_vazia = (not tem_docs_contrib) and contrib.receita_f550 <= 0 and bool(icms.docs) \
        and not r.contrib_ausente

    r.indice = indice
    return r


# ============================================================================
# 4. PLANILHA DECISÓRIA (padrão AgriTax)
# ============================================================================

VERDE, OURO, CHUMBO, AZUL = "7AB82E", "F2B705", "404040", "0000CC"


def gerar_decisoria(resultados, cnpj, saida):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    f_hdr = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    f_txt = Font(name="Arial", size=10, color=CHUMBO)
    f_azul = Font(name="Arial", size=10, color=AZUL)
    f_fech = Font(name="Arial", size=10, color=CHUMBO, italic=True)
    fill_v = PatternFill("solid", fgColor=VERDE)
    fill_o = PatternFill("solid", fgColor=OURO)
    wrap = Alignment(wrap_text=True, vertical="top")

    def cab(ws, cols, larguras):
        for j, (titulo, larg) in enumerate(zip(cols, larguras), start=1):
            cel = ws.cell(row=1, column=j, value=titulo)
            cel.font, cel.fill, cel.alignment = f_hdr, fill_v, wrap
            ws.column_dimensions[get_column_letter(j)].width = larg
        ws.freeze_panes = "A2"

    def dv(ws, col_letra, opcoes, n_linhas):
        d = DataValidation(type="list", formula1='"%s"' % ",".join(opcoes), allow_blank=True)
        ws.add_data_validation(d)
        d.add("%s2:%s%d" % (col_letra, col_letra, max(n_linhas + 1, 2)))

    comp_txt = lambda c: "%02d/%d" % (c[1], c[0]) if c else "(definir período aberto)"

    def escreve(ws, lin, valores, fonte, fmt_moeda=()):
        for j, v in enumerate(valores, start=1):
            cel = ws.cell(row=lin, column=j, value=v)
            cel.font, cel.alignment = fonte, wrap
            if j in fmt_moeda:
                cel.number_format = '#,##0.00'

    # ------------------------------------------------------------ Instruções
    ws = wb.active
    ws.title = "Instruções"
    ws.column_dimensions["A"].width = 112
    linhas = [
        ("AgriTax Audit — Módulo Frigorífico | Planilha Decisória", True),
        ("CNPJ analisado: %s" % cnpj, False),
        ("", False),
        ("COMO USAR: revise cada linha, escolha SUA DECISÃO no dropdown e registre a JUSTIFICATIVA/PROVA.", False),
        ("Linhas em AZUL vieram do cruzamento EFD-ICMS × EFD-Contribuições (a nota existe no ICMS e não nas Contribuições) — comprovam a ORIGEM do crédito.", False),
        ("", False),
        ("REGRA DE PERÍODOS FECHADOS (PER/DCOMP DE RESSARCIMENTO):", True),
        ("O trabalho parte de TODOS os pedidos de ressarcimento já efetuados + consulta de situação. Período com PER DEFERIDA, ANÁLISE CONCLUÍDA ou HOMOLOGADO = FECHADO para inserção/retificação.", False),
        ("• OPORTUNIDADE em período fechado: é mapeada e trazida mesmo assim. A apropriação extemporânea no primeiro período aberto (coluna COMPETÊNCIA DE DESTINO) vale SOMENTE para benefícios inseríveis por DOCUMENTO — registros dos BLOCOS A (serviços), C (NF-e), D (transportes) e F (demais operações/imobilizado) da EFD-Contribuições. Ver coluna REGISTRO DE INSERÇÃO.", False),
        ("• OPORTUNIDADE em período fechado que só se materializa como ajuste de apuração (bloco M): trazida com ação SEM INSERÇÃO EM PERÍODO ABERTO — avaliar via alternativa (não entra no EFD de destino).", False),
        ("• RISCO em período fechado: não vira ação de correção — compõe a aba MATRIZ DE RISCO (exposição a gerir: provisionar, monitorar, preparar defesa).", False),
        ("• RISCO em período aberto: segue em Expurgos-Riscos como ação de correção (estornar/retificar).", False),
        ("", False),
        ("DICIONÁRIO DE AÇÕES:", True),
        ("CREDITAR — inserir o crédito na EFD-Contribuições da competência (retificadora).", False),
        ("APROPRIAR EXTEMPORÂNEO — inserir o crédito de período fechado na EFD do período de destino (aberto), com memória de origem.", False),
        ("INSERIR NA EFD — escriturar documento/registro ausente.", False),
        ("ESTORNAR — reverter crédito indevido (risco em período aberto).", False),
        ("RECLASSIFICAR CST — corrigir o CST da operação (ex.: 06 → 09 na venda p/ indústria).", False),
        ("RETIFICAR COMPETÊNCIA — reabrir a competência de origem (apenas períodos abertos).", False),
        ("TRANSMITIR PER/DCOMP — pedir ressarcimento/compensação da parcela vinculada à exportação.", False),
        ("AGUARDAR PROVA — fruição condicionada a documento externo (GTA, DU-E, laudo, CNAE do adquirente).", False),
        ("PROVISIONAR / MONITORAR / PREPARAR DEFESA — ações de gestão da Matriz de Risco.", False),
        ("", False),
        ("BASE LEGAL CENTRAL: arts. 32-37 da Lei 12.058/2009; SC COSIT 80/2021 e 188/2021; Leis 10.637/2002 e 10.833/2003; LC 224/2025 (efeitos a partir de abr/2026); Lei 14.592/2023 (expurgo do ICMS).", False),
    ]
    for i, (t, b) in enumerate(linhas, start=1):
        cel = ws.cell(row=i, column=1, value=t)
        cel.font = Font(name="Arial", bold=b, size=11 if b else 10, color=CHUMBO)
        cel.alignment = wrap
        if i == 1:
            cel.fill = fill_o

    # ---------------------------------------------------------------- Resumo
    ws = wb.create_sheet("Resumo")
    cab(ws, ["Competência", "Situação PER/DCOMP", "Período", "Receita bruta",
             "Receita export. (lista art.33)", "Índice §7º",
             "Base gado (PF+PJ susp.)", "Presumido devido", "Presumido escriturado",
             "Diferença (oportunidade)", "Créditos descobertos B1/B2",
             "Competência de destino", "Prescrição (data-limite)",
             "Fonte da receita", "Apuração (bloco M)", "Regime PIS/COFINS (0110)",
             "Estabelecimentos (EFD-ICMS)"],
        [12, 20, 11, 16, 18, 10, 18, 15, 15, 17, 17, 16, 15, 24, 20, 22, 26])
    lin = 2
    for r in resultados:
        limite = date(r.comp[0] + 5, r.comp[1], 1)
        cruz = sum(x["credito"] for x in r.b1) + sum(x["credito"] for x in r.b2)
        vals = [comp_txt(r.comp), r.per_status,
                "FECHADO" if r.fechado else "ABERTO",
                r.receita_total, r.receita_export_lista, r.indice,
                r.base_gado_pf + r.base_gado_pj, r.presumido_devido,
                r.presumido_escriturado,
                max(r.presumido_devido - r.presumido_escriturado, 0.0),
                cruz,
                comp_txt(r.destino) if r.fechado else "(o próprio)",
                limite.strftime("%d/%m/%Y"), r.fonte_receita,
                ("SEM EFD-CONTRIBUIÇÕES NO ACERVO" if r.contrib_ausente else
                 ("SEM MOVIMENTO" if r.contrib_vazia else
                  ("AUSENTE (M vazio)" if r.sem_apuracao else
                   ("CUMULATIVO (F550, sem bloco M)" if r.cumulativo else "OK")))),
                {"1": "NÃO CUMULATIVO", "2": "CUMULATIVO", "3": "AMBOS"}.get(r.regime, r.regime or "?"),
                ", ".join(r.estabs) if r.estabs else ""]
        for j, v in enumerate(vals, start=1):
            cel = ws.cell(row=lin, column=j, value=v)
            cel.font = f_fech if r.fechado and j <= 3 else f_txt
            if j in (4, 5, 7, 8, 9, 10, 11):
                cel.number_format = '#,##0.00'
            if j == 6:
                cel.number_format = '0.0%'
        lin += 1

    # ------------------------------------------------------- Presumido_Art33
    ws = wb.create_sheet("Presumido_Art33")
    cols = ["Competência", "Situação PER/DCOMP", "Base gado PF",
            "Base gado PJ c/ suspensão", "Índice export. §7º",
            "Alíquota vigente (PIS+COFINS)", "Presumido devido",
            "Presumido escriturado", "Oportunidade", "Competência de destino",
            "Legislação", "Exporta? (SIM/NÃO)", "AÇÃO RESULTANTE",
            "SUA DECISÃO", "JUSTIFICATIVA/PROVA"]
    cab(ws, cols, [12, 18, 14, 17, 11, 13, 14, 14, 14, 16, 44, 12, 30, 20, 32])
    lin = 2
    for r in resultados:
        ap, ac = aliq_presumido(r.comp)
        leg = ("Art. 33 e §7º, Lei 12.058/2009; SC COSIT 80/2021 e 188/2021; "
               "LC 224/2025 (45% a partir de 04/2026); material seções 2/8.2/9")
        fonte = f_fech if r.fechado else f_txt
        ws.cell(row=lin, column=1, value=comp_txt(r.comp)).font = fonte
        ws.cell(row=lin, column=2, value=r.per_status).font = fonte
        for j, v in ((3, r.base_gado_pf), (4, r.base_gado_pj), (5, r.indice),
                     (6, ap + ac), (7, r.presumido_devido),
                     (8, r.presumido_escriturado),
                     (9, max(r.presumido_devido - r.presumido_escriturado, 0.0))):
            cel = ws.cell(row=lin, column=j, value=v)
            cel.font = fonte
            cel.number_format = '0.000%' if j in (5, 6) else '#,##0.00'
        ws.cell(row=lin, column=10,
                value=comp_txt(r.destino) if r.fechado else "(o próprio)").font = fonte
        ws.cell(row=lin, column=11, value=leg).font = fonte
        ws.cell(row=lin, column=11).alignment = wrap
        ws.cell(row=lin, column=12, value="SIM" if r.indice > 0 else "NÃO").font = fonte
        if r.fechado:
            acao_sim = ("APROPRIAR EXTEMPORÂNEO EM %s — bloco C (C100/C170, CST 61-66, "
                        "docs de aquisição do gado) + reflexo M100/M105 + TRANSMITIR PER/DCOMP"
                        % comp_txt(r.destino))
        else:
            acao_sim = "RETIFICAR COMPETÊNCIA + TRANSMITIR PER/DCOMP"
        ws.cell(row=lin, column=13,
                value='=IF(L%d="SIM","%s","AGUARDAR PROVA")' % (lin, acao_sim)).font = fonte
        ws.cell(row=lin, column=13).alignment = wrap
        lin += 1
    n = lin - 2
    dv(ws, "L", ["SIM", "NÃO"], n)
    dv(ws, "N", ["ACEITAR", "VALIDAR", "CORRIGIR", "REVERTER"], n)

    # --------------------------------------------------------------- Créditos
    ws = wb.create_sheet("Créditos")
    cols = ["Competência origem", "Situação PER/DCOMP", "Competência de destino",
            "Registro de inserção (bloco)", "Classe de insumo", "Fornecedor",
            "Chave NF-e/CT-e", "Valor da operação", "Crédito potencial (9,25%)",
            "Origem do documento", "Legislação", "AÇÃO EXECUTIVA",
            "SUA DECISÃO", "JUSTIFICATIVA/PROVA"]
    cab(ws, cols, [13, 18, 15, 18, 17, 22, 44, 14, 15, 25, 36, 26, 16, 26])
    lin = 2
    for r in resultados:
        destino = comp_txt(r.destino) if r.fechado else "(o próprio)"
        acao_nf = ("APROPRIAR EXTEMPORÂNEO EM %s" % comp_txt(r.destino)) if r.fechado \
            else "CREDITAR + INSERIR NA EFD"
        if r.cumulativo:
            # Lucro Presumido/cumulativo: art. 8º, II, Lei 10.637 e art. 10, II, Lei 10.833
            acao_nf = "SEM DIREITO A CRÉDITO (regime cumulativo) — SIMULAÇÃO p/ Lucro Real"
            destino = "(n/a — cumulativo)"
        elif r.contrib_ausente:
            acao_nf = "UNIVERSO CREDITÁVEL — CONFERIR NA EFD-CONTRIBUIÇÕES (arquivo ausente)"
        for x in r.b1:
            escreve(ws, lin, [comp_txt(r.comp), r.per_status, destino,
                              "C — C100/C170", x["classe"],
                              x["fornecedor"], x["chave"], x["valor"], x["credito"],
                              ("EFD-ICMS (universo — Contribuições não fornecida)" if r.contrib_ausente
                               else "Cruzamento EFD-ICMS (ausente nas Contribuições)"),
                              "Art. 3º, II, Leis 10.637/02 e 10.833/03; PN COSIT 5/2018; manutenção p/ vendas a alíq. zero: art. 17 Lei 11.033/04 e art. 16 Lei 11.116/05; material 6.1",
                              acao_nf], f_azul, fmt_moeda=(8, 9))
            lin += 1
        for x in r.b2:
            escreve(ws, lin, [comp_txt(r.comp), r.per_status, destino,
                              "D — D100", "Frete (CT-e)",
                              "", x["chave"], x["vl_serv"], x["credito"],
                              "Cruzamento EFD-ICMS D100 (ausente nas Contribuições)",
                              "Art. 3º, II e IX, Leis 10.637/02 e 10.833/03; material 6.1-f/g",
                              acao_nf], f_azul, fmt_moeda=(8, 9))
            lin += 1
        for cl, (total, tem) in sorted(r.b3_classes.items()):
            if tem:
                continue
            escreve(ws, lin, [comp_txt(r.comp), r.per_status, destino,
                              "A/C — docs da classe", cl,
                              "(classe inteira)", "", total, total * sum(ALIQ_BASICA),
                              "B3 — classe sem NENHUM crédito no período",
                              "Tema 779/STJ; PN COSIT 5/2018; material 6.1",
                              acao_nf if r.cumulativo else
                              (("APROPRIAR EXTEMPORÂNEO EM %s" % comp_txt(r.destino)) if r.fechado else "CREDITAR")],
                    f_txt, fmt_moeda=(8, 9))
            lin += 1
    dv(ws, "M", ["CREDITAR", "EXCLUIR", "MANTER"], max(lin - 2, 1))

    # ----------------------------------------------------------- Receitas-CST
    ws = wb.create_sheet("Receitas-CST")
    cols = ["Competência", "Situação PER/DCOMP", "Achado", "Detalhe", "Valor",
            "Legislação", "AÇÃO EXECUTIVA", "SUA DECISÃO", "JUSTIFICATIVA/PROVA"]
    cab(ws, cols, [12, 18, 32, 38, 14, 40, 26, 20, 28])
    lin = 2
    for r in resultados:
        fonte = f_fech if r.fechado else f_txt
        if r.receita_export_fora > 0:
            if r.fechado:
                acao = ("SEM INSERÇÃO EM PERÍODO ABERTO (ajuste de apuração — bloco M; "
                        "período fechado): aplicar apenas no recálculo do índice do extemporâneo")
            else:
                acao = "RECLASSIFICAR (recalcular índice)"
            escreve(ws, lin, [comp_txt(r.comp), r.per_status,
                              "A5 — exportação de NCM fora da lista (couro/tripas) no rateio",
                              "Excluir do numerador do índice §7º", r.receita_export_fora,
                              "Art. 33, caput (lista exaustiva); SC COSIT 80/2021",
                              acao], fonte, fmt_moeda=(5,))
            lin += 1
        for adq, val in sorted(r.a6_por_adquirente.items(), key=lambda kv: -kv[1])[:20]:
            escreve(ws, lin, [comp_txt(r.comp), r.per_status,
                              "A6 — venda de carne p/ CNPJ: conferir suspensão (art. 32, II)",
                              "Adquirente: %s — industrializa? Se SIM, CST 09" % adq, val,
                              "Art. 32, II, Lei 12.058/2009; material 6.1-i",
                              "AGUARDAR PROVA (CNAE do adquirente) → RECLASSIFICAR CST (prospectivo)"],
                    fonte, fmt_moeda=(5,))
            lin += 1
    dv(ws, "H", ["MANTER", "RECLASSIFICAR CST", "TRIBUTAR", "AGUARDAR"], max(lin - 2, 1))

    # -------------------------------------------- Expurgos-Riscos (abertos)
    ws = wb.create_sheet("Expurgos-Riscos")
    cols = ["Competência", "Risco", "Detalhe", "Valor exposto", "Legislação",
            "AÇÃO EXECUTIVA", "SUA DECISÃO", "JUSTIFICATIVA/PROVA"]
    cab(ws, cols, [12, 38, 42, 15, 42, 22, 18, 28])
    lin = 2
    matriz = []   # riscos de períodos fechados → Matriz de Risco
    for r in resultados:
        itens = []
        for x in r.risco_a4:
            itens.append(("A4 — crédito básico sobre gado de PJ (compra é SUSPENSA)",
                          "Chave %s" % x["chave"], x["credito"],
                          "Art. 32, Lei 12.058/2009; art. 3º, §2º, II, Leis 10.637/10.833; SC COSIT 80/2021",
                          "ESTORNAR"))
        if r.risco_a7 > 0:
            ap10, ac10 = aliq_carne_interna(r.comp)
            itens.append(("A7 — carnes internas com CST 06 após LC 224 (devido 10% da alíquota)",
                          "Débito não apurado: base × %.3f%%" % ((ap10 + ac10) * 100),
                          r.risco_a7 * (ap10 + ac10),
                          "LC 224/2025; IN RFB 2.305/2025; material 3.2",
                          "TRIBUTAR (retificar)"))
        for x in r.risco_b4[:50]:
            itens.append(("B4 — base de crédito sem expurgo do ICMS (>= mai/2023)",
                          "Chave %s — BC=%.2f = valor do item" % (x["chave"], x["bc"]),
                          x["bc"], "Lei 14.592/2023; material 6.4.1", "CORRIGIR BASE"))
        if r.presumido_escriturado > 0 and r.indice == 0:
            itens.append(("A8 — crédito PRESUMIDO escriturado SEM receita de exportação no mês",
                          "Art. 33 exige produção destinada à exportação (posição RFB/CARF) — presumido sem lastro",
                          r.presumido_escriturado,
                          "Art. 33, Lei 12.058/2009; SC COSIT 80/2021 e 188/2021; material 2/8.3",
                          "ESTORNAR / AVALIAR DEFESA"))
        if r.sem_apuracao:
            itens.append(("B5 — documentos escriturados SEM apuração (bloco M vazio)",
                          "C100/C170 presentes e nenhum M100/M500: créditos e débitos do mês não apurados",
                          0,
                          "IN RFB 2.121/2022 (escrituração); art. 3º Leis 10.637/10.833; material 9",
                          "RETIFICAR (apurar bloco M)"))
        if r.contrib_vazia:
            itens.append(("B5 — EFD-Contribuições SEM MOVIMENTO com EFD-ICMS movimentada",
                          "Obrigação entregue zerada em mês com operações no ICMS",
                          0,
                          "IN RFB 2.121/2022; multa art. 12 Lei 8.218/91; material 9",
                          "RETIFICAR (escriturar o mês)"))
        for trib, efetiva, esperada in r.a3_alertas:
            itens.append(("A3 — alíquota do presumido fora da vigência",
                          "%s: efetiva %.4f%% × esperada %.4f%%" % (trib, efetiva * 100, esperada * 100),
                          0, "Art. 33, Lei 12.058/2009; LC 224/2025", "CORRIGIR BASE"))
        for it in itens:
            if r.fechado:
                matriz.append((r,) + it)
            else:
                escreve(ws, lin, [comp_txt(r.comp)] + list(it[:4]) + [it[4]],
                        f_txt, fmt_moeda=(4,))
                lin += 1
    dv(ws, "G", ["ESTORNAR", "CORRIGIR", "MANTER", "AGUARDAR"], max(lin - 2, 1))

    # ---------------------------------------------------- Matriz de Risco
    ws = wb.create_sheet("Matriz de Risco")
    cols = ["Competência origem", "Situação PER/DCOMP", "Risco", "Detalhe",
            "Exposição (R$)", "Por que não é corrigível",
            "Legislação", "AÇÃO DE GESTÃO", "SUA DECISÃO", "JUSTIFICATIVA/PROVA"]
    cab(ws, cols, [14, 20, 36, 38, 15, 34, 38, 20, 18, 28])
    lin = 2
    for r, risco, det, valor, leg, _acao in matriz:
        escreve(ws, lin, [comp_txt(r.comp), r.per_status, risco, det, valor,
                          "Período com ressarcimento decidido (%s): retificação inviável — gerir a exposição"
                          % r.per_status, leg, "PROVISIONAR"], f_txt, fmt_moeda=(5,))
        lin += 1
    if lin == 2:
        ws.cell(row=2, column=1, value="(sem riscos em períodos fechados)").font = f_txt
    dv(ws, "I", ["PROVISIONAR", "MONITORAR", "PREPARAR DEFESA", "ACEITAR RISCO"], max(lin - 2, 1))

    wb.save(saida)
    return saida


# ============================================================================
# 5. ORQUESTRAÇÃO / CLI
# ============================================================================

def _parse_desde(txt):
    """'MM/AAAA' -> (ano, mes) ou None."""
    if not txt:
        return None
    m = re.match(r"^\s*(\d{1,2})/(\d{4})\s*$", txt)
    if not m:
        raise SystemExit("--desde deve ser MM/AAAA (ex.: 09/2021)")
    return (int(m.group(2)), int(m.group(1)))


def executar(pasta, saida, per_csv=None, desde=None, sem_contrib=False):
    situacao = carregar_situacao_per(per_csv)
    desde = _parse_desde(desde) if isinstance(desde, str) else desde
    if per_csv and not situacao:
        print("AVISO: --per informado mas nenhuma competência lida de %s" % per_csv)
    # comp -> {"ICMS": {cnpj: (prio, path)}, "CONTRIB": (prio, path)}
    arquivos = defaultdict(lambda: {"ICMS": {}, "CONTRIB": None})
    cnpjs = set()
    for nome in sorted(os.listdir(pasta)):
        if not nome.lower().endswith(".txt"):
            continue
        p = os.path.join(pasta, nome)
        tipo, comp, cnpj, prio = detectar_tipo(p)
        if tipo and comp and desde and comp < desde:
            continue                                   # fora do escopo temporal
        if tipo and comp:
            if tipo == "ICMS":
                atual = arquivos[comp]["ICMS"].get(cnpj)
                if atual is None or prio > atual[0]:
                    arquivos[comp]["ICMS"][cnpj] = (prio, p)
            else:
                atual = arquivos[comp]["CONTRIB"]
                if atual is None or prio > atual[0]:
                    arquivos[comp]["CONTRIB"] = (prio, p)
            if cnpj:
                cnpjs.add(cnpj)
    for comp in sorted(arquivos):
        for cnpj, (prio, p) in sorted(arquivos[comp]["ICMS"].items()):
            print("  [ICMS %02d/%d %s]%s %s" % (comp[1], comp[0], cnpj[-6:],
                  " (subst)" if prio and prio[0] else "", os.path.basename(p)[:70]))
        if arquivos[comp]["CONTRIB"]:
            prio, p = arquivos[comp]["CONTRIB"]
            print("  [CONTRIB %02d/%d]%s %s" % (comp[1], comp[0],
                  " (retif)" if prio and prio[0] else "", os.path.basename(p)[:70]))
    if len(cnpjs) > 1:
        print("ATENÇÃO: %d CNPJs na pasta (matriz+filiais consolidados por competência): %s"
              % (len(cnpjs), ", ".join(sorted(cnpjs))))
    n_sem_contrib = sum(1 for c in arquivos if arquivos[c]["ICMS"] and not arquivos[c]["CONTRIB"])
    if n_sem_contrib and not sem_contrib:
        print("AVISO: %d competência(s) com EFD-ICMS e SEM EFD-Contribuições — rodando em modo "
              "'universo ICMS' (B1/B2 = tudo o que é creditável; conferir quando a Contribuições chegar)"
              % n_sem_contrib)
    resultados = []
    for comp in sorted(arquivos):
        par = arquivos[comp]
        if not par["ICMS"]:
            print("  ! Competência %02d/%d sem EFD-ICMS — pulando" % (comp[1], comp[0]))
            continue
        paths_icms = [p for _, p in sorted(par["ICMS"].values(), key=lambda x: x[1])]
        icms = EfdIcmsConsolidada(paths_icms) if len(paths_icms) > 1 else EfdIcms(paths_icms[0])
        if not hasattr(icms, "cnpjs"):
            icms.cnpjs = list(par["ICMS"].keys())
        contrib = EfdContrib(par["CONTRIB"][1]) if par["CONTRIB"] and not sem_contrib else EfdContribVazia()
        r = revisar_competencia(comp, icms, contrib)
        resultados.append(r)
        print("  %02d/%d | receita %.2f | índice %.1f%% | presumido devido %.2f × escriturado %.2f | B1/B2: %d achados"
              % (comp[1], comp[0], r.receita_total, r.indice * 100,
                 r.presumido_devido, r.presumido_escriturado, len(r.b1) + len(r.b2)))
    if not resultados:
        print("Nenhuma competência completa para revisar.")
        return None
    # ---- Regra de períodos fechados (PER de ressarcimento decidido) --------
    abertas = [r.comp for r in resultados
               if not periodo_fechado(situacao.get(r.comp, ""))]
    for r in resultados:
        st = situacao.get(r.comp, "")
        r.per_status = st if st else "SEM PER / ABERTO"
        r.fechado = periodo_fechado(st)
        if r.fechado:
            # primeiro período ABERTO >= origem; senão o primeiro aberto geral;
            # senão fica em aberto p/ o analista definir (período corrente)
            depois = [c for c in abertas if c >= r.comp]
            r.destino = (depois[0] if depois else (abertas[0] if abertas else None))
            print("  * %02d/%d FECHADO (%s) -> oportunidades remapeadas p/ %s"
                  % (r.comp[1], r.comp[0], st,
                     "%02d/%d" % (r.destino[1], r.destino[0]) if r.destino
                     else "período aberto a definir"))
    caminho = gerar_decisoria(resultados, "/".join(sorted(cnpjs)) or "?", saida)
    print("Planilha decisória gerada: %s" % caminho)
    return caminho


def main():
    ap = argparse.ArgumentParser(description="AgriTax Audit — Módulo Frigorífico (Fase 1)")
    ap.add_argument("pasta", nargs="?", help="Pasta com os .txt das EFDs")
    ap.add_argument("--saida", default="decisoria_frigorifico.xlsx")
    ap.add_argument("--per", default=None,
                    help="CSV 'competencia;situacao' dos PERs de ressarcimento (ex.: 05/2026;PER DEFERIDA)")
    ap.add_argument("--desde", default=None,
                    help="Primeira competência do escopo, MM/AAAA (prescrição corre mês a mês)")
    ap.add_argument("--sem-contrib", action="store_true",
                    help="Ignorar EFD-Contribuições mesmo se existir (modo universo ICMS)")
    args = ap.parse_args()
    pasta = args.pasta
    if not pasta:
        try:
            import tkinter as tk
            from tkinter import filedialog
            root = tk.Tk(); root.withdraw()
            pasta = filedialog.askdirectory(title="Pasta com as EFDs do frigorífico")
            root.destroy()
        except Exception:
            pass
    if not pasta or not os.path.isdir(pasta):
        print("Informe a pasta com os arquivos .txt das EFDs.")
        sys.exit(1)
    executar(pasta, args.saida, args.per, args.desde, args.sem_contrib)


if __name__ == "__main__":
    main()
